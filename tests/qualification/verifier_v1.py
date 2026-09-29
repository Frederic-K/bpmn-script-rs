"""Attentes V1 des 30 scénarios du pilote de qualification.

Le paquet docs/dossier-claude/references/qualification-rust-reproductible.zip
compare le binaire Rust au programme Python d'origine. Son check_qualification.py
décrit le comportement historique ; ce script décrit le contrat V1
(docs/m3-contrats-v1.md). Aucune donnée de référence n'est modifiée.

    python qualify_rust.py --root <paquet décompressé> --binary <bpmn-script-rs>
    python tests/qualification/verifier_v1.py <paquet décompressé>/rust-qualification-results.json

Les sorties de chaque scénario sont relues dans <paquet décompressé>/qualification-runs/
(dernière exécution du scénario) pour contrôler l'évolution de la colonne F.
"""

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

# Code de sortie V1 : 0 étape terminée, 1 erreur bloquante, 2 décisions fournies sans SGX.
CODES = {
    0: {"inventory", "dry_run", "accepted", "mixed", "normalized_yes", "empty_mapping",
        "unknown_mapping", "duplicate_mapping", "missing_output", "partial", "same_name",
        "zip_attributes"},
    1: {"no_source", "multiple_sources", "invalid_zip", "numeric_old", "numeric_new",
        "formula_new", "null_name", "bad_children"},
    2: {"refused", "pending", "stale_name", "stale_count", "stale_path", "stale_flux",
        "duplicate_validation", "chain_conflict", "duplicate_different", "old_output"},
}

PRODUIT_UN_SGX = {"accepted", "mixed", "normalized_yes", "partial", "same_name", "zip_attributes"}

# Écarts de contenu avec Python, tous volontaires (P02, P04, P05) ; tout autre
# écart est une régression.
ECARTS_AVEC_PYTHON = {
    "empty_mapping": "P05 : analyse vide, Python échoue",
    "unknown_mapping": "P05 : analyse vide, Python échoue",
    "missing_output": "P05 : dossier de sortie créé, Python échoue",
    "numeric_old": "P02 : classeur refusé",
    "numeric_new": "P02 : classeur refusé",
    "formula_new": "P02 : classeur refusé",
    "pending": "P02 : motif « valeur non reconnue » explicite",
}

resultats = {resultat["scenario"]: resultat for resultat in json.load(open(sys.argv[1], encoding="utf-8"))}
EXECUTIONS = Path(sys.argv[1]).resolve().parent / "qualification-runs"
echecs = []


def verifier(condition, message):
    if not condition:
        echecs.append(message)


# Évolution V1 de la colonne F (docs/m3-contrats-v1.md) : seuls ces deux
# fichiers peuvent différer de Python, et seulement de la façon attendue.
# Leur écart est retiré de la liste après un contrôle exact ; un écart
# différent reste une régression.
ANCIEN_ENTETE_F = "Flux avec occurrences multiples"
NOUVEL_ENTETE_F = "Répétitions dans un même modèle"


def dossier_execution(nom):
    # qualify_rust.py numérote les réexécutions : nom, nom-1, nom-2…
    candidats = [d for d in EXECUTIONS.iterdir() if re.fullmatch(rf"{re.escape(nom)}(-\d+)?", d.name)]
    return max(candidats, key=lambda d: int(d.name[len(nom) + 1:] or 0))


def colonne_f_attendue(par_modele):
    repetitions = [
        f"{modele['flux']} ({modele['occurrences']} occurrences) — {chemin}"
        for chemin, modele in par_modele.items()
        if modele["occurrences"] > 1
    ]
    return "\n".join(repetitions) or None


def ecart_synthese_conforme(nom, python, rust):
    synthese_python = json.loads((python / "synthese.json").read_text("utf-8"))
    synthese_rust = json.loads((rust / "synthese.json").read_text("utf-8"))
    conforme = list(synthese_python) == list(synthese_rust)
    for lane, infos in synthese_python.items():
        attendu = {cle: valeur for cle, valeur in infos.items() if cle != "occurrences_par_flux"}
        obtenu = synthese_rust.get(lane)
        conforme &= "occurrences_par_flux" in infos
        conforme &= obtenu == attendu and list(obtenu) == list(attendu)
    verifier(conforme, f"{nom} : synthese.json diffère au-delà du retrait de occurrences_par_flux")
    return conforme


def ecart_inventaire_conforme(nom, python, rust):
    synthese_python = json.loads((python / "synthese.json").read_text("utf-8"))
    classeurs = [load_workbook(dossier / "inventaire_swimlanes.xlsx") for dossier in (python, rust)]
    feuilles = [[feuille.title for feuille in classeur] for classeur in classeurs]
    lignes_python, lignes_rust = (
        [list(ligne) for ligne in classeur["Correspondance"].iter_rows(values_only=True)]
        for classeur in classeurs
    )
    conforme = feuilles == [["Correspondance"], ["Correspondance"]]
    conforme &= len(lignes_python) == len(lignes_rust)
    conforme &= all(len(ligne) == 6 for ligne in lignes_python + lignes_rust)
    if conforme:
        conforme &= lignes_python[0] == lignes_rust[0][:5] + [ANCIEN_ENTETE_F]
        conforme &= lignes_rust[0][5] == NOUVEL_ENTETE_F
        for avant, apres in zip(lignes_python[1:], lignes_rust[1:]):
            conforme &= avant[:5] == apres[:5]
            conforme &= apres[5] == colonne_f_attendue(synthese_python.get(apres[0], {}).get("occurrences_par_modele", {}))
    verifier(conforme, f"{nom} : inventaire_swimlanes.xlsx diffère au-delà de la colonne F attendue")
    return conforme


def ecarts_hors_colonne_f(nom, ecarts):
    controles = {
        "synthese.json": ecart_synthese_conforme,
        "inventaire_swimlanes.xlsx": ecart_inventaire_conforme,
    }
    restants = []
    for ecart in ecarts:
        controle = controles.get(ecart["file"])
        if ecart["reason"] == "content" and controle:
            dossier = dossier_execution(nom)
            if controle(nom, dossier / "python/output", dossier / "rust/output"):
                continue
        restants.append(ecart)
    return restants


attendus = set().union(*CODES.values())
verifier(attendus <= resultats.keys(), f"scénarios manquants : {sorted(attendus - resultats.keys())}")

for code, scenarios in CODES.items():
    for nom in scenarios & resultats.keys():
        rust = resultats[nom]["rust"]
        verifier(rust["code"] == code, f"{nom} : code {rust['code']}, attendu {code}")
        verifier(rust["source_intact"], f"{nom} : source modifiée")
        nouveau_sgx = bool(rust["sgx_after"]) and rust["sgx_after"] != rust["sgx_before"]
        verifier(nouveau_sgx == (nom in PRODUIT_UN_SGX), f"{nom} : SGX produit = {nouveau_sgx}")
        ecarts = ecarts_hors_colonne_f(nom, resultats[nom]["semantic_differences"])
        ecart = bool(ecarts)
        verifier(ecart == (nom in ECARTS_AVEC_PYTHON), f"{nom} : écart avec Python = {ecart} {ecarts}")

# L'évolution de la colonne F doit être effective : la fixture de référence
# répète « A » dans un modèle homonyme d'un autre.
verifier(
    {"inventaire_swimlanes.xlsx", "synthese.json"}
    <= {ecart["file"] for ecart in resultats["inventory"]["semantic_differences"]},
    "inventory : colonne F ou synthese.json identiques à Python, évolution absente",
)

# Essai facultatif sur une copie de SAPHIR (docs/m6-qualification.md) : à
# chaque étape, seuls les deux fichiers de la colonne F peuvent différer.
if "saphir" in resultats:
    for etape in resultats["saphir"]["stages"]:
        autres = set(etape["differences"]) - {"inventaire_swimlanes.xlsx", "synthese.json"}
        verifier(not autres, f"saphir, étape {etape['stage']} : écart avec Python {sorted(autres)}")
    dossier = dossier_execution("saphir")
    ecart_synthese_conforme("saphir", dossier / "python/output", dossier / "rust/output")
    ecart_inventaire_conforme("saphir", dossier / "python/output", dossier / "rust/output")

controle = resultats["zip_attributes"]["rust"]["archive_check"]
verifier(controle["untouched_contents_preserved"], "zip_attributes : entrées non modifiées altérées")
verifier(controle["archive_comment_preserved"], "zip_attributes : commentaire d'archive perdu (P06)")
entree_modifiee = next(ecart for ecart in controle["attribute_changes"] if ecart["entry"] == "dossier/a/model_1_.json")
verifier(
    not {"extra", "comment", "date_time", "compress_type"} & set(entree_modifiee["changed_attributes"]),
    f"zip_attributes : attributs perdus sur l'entrée modifiée {entree_modifiee['changed_attributes']} (P06)",
)
verifier(
    resultats["old_output"]["rust"]["sgx_before"] == resultats["old_output"]["rust"]["sgx_after"],
    "old_output : ancienne sortie modifiée",
)

if echecs:
    print("Attentes V1 non respectées :")
    for echec in echecs:
        print(" -", echec)
    sys.exit(1)
print(f"Attentes V1 respectées : {len(attendus)} scénarios.")
