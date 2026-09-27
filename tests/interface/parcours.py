"""Test de bout en bout de l'application : vraie fenêtre, vrai moteur Rust.

Pilote l'application compilée par WebDriver (tauri-driver), remplit les vrais
dialogues de fichiers au clavier (xdotool) et modifie les classeurs comme le
ferait Excel (openpyxl). Enregistre une capture de chaque écran.

Prérequis (Linux) : tauri-driver, WebKitWebDriver, Xvfb, xdotool, Python avec
selenium et openpyxl.

    xvfb-run -a python tests/interface/parcours.py \\
        --application target/debug/bpmn-script-app --captures captures/
"""

import argparse
import json
import shutil
import subprocess
import tempfile
import time
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from selenium import webdriver
from selenium.common.exceptions import NoSuchElementException, StaleElementReferenceException, WebDriverException
from selenium.webdriver.common.by import By
from selenium.webdriver.common.options import ArgOptions
from selenium.webdriver.support.ui import WebDriverWait

MODELE_A = "a/model_1_.json"
MODELE_B = "b/model_1_.json"

parametres = argparse.ArgumentParser()
parametres.add_argument("--application", type=Path, required=True)
parametres.add_argument("--captures", type=Path, required=True)
parametres.add_argument("--tauri-driver", default="tauri-driver")
arguments = parametres.parse_args()
arguments.captures.mkdir(parents=True, exist_ok=True)


def lane(nom):
    return {"stencil": {"id": "Lane"}, "properties": {"name": nom}, "childShapes": []}


def ecrire_sgx(chemin):
    modeles = {
        MODELE_A: {"childShapes": [lane("Serv. achats"), lane("Serv. achats"), lane("Direction")]},
        MODELE_B: {"childShapes": [lane("Serv. achats")]},
    }
    with zipfile.ZipFile(chemin, "w", zipfile.ZIP_DEFLATED) as archive:
        for nom, modele in modeles.items():
            archive.writestr(nom, json.dumps(modele, ensure_ascii=False))
            archive.writestr(nom.replace("model_1_", "model_meta"), json.dumps({"name": "Achats"}))


def ecrire_classeur(chemin, feuille, lignes):
    classeur = Workbook()
    classeur.active.title = feuille
    for ligne in lignes:
        classeur.active.append(ligne)
    classeur.save(chemin)


def lanes_du_sgx(chemin, modele):
    with zipfile.ZipFile(chemin) as archive:
        return [forme["properties"]["name"] for forme in json.loads(archive.read(modele))["childShapes"]]


# ---------------------------------------------------------------- Pilotage

espace = Path(tempfile.mkdtemp(prefix="bpmn-interface-"))
source = espace / "export.sgx"
parent = espace / "traitements"
parent.mkdir()
ecrire_sgx(source)
dossier = parent / "export - traitement"

pilote = subprocess.Popen([arguments.tauri_driver], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
time.sleep(2)
options = ArgOptions()
options.set_capability("browserName", "wry")
options.set_capability("tauri:options", {"application": str(arguments.application.resolve())})
fenetre = webdriver.Remote(command_executor="http://127.0.0.1:4444", options=options)
fenetre.set_window_size(1200, 800)
attente = WebDriverWait(fenetre, 15, ignored_exceptions=[StaleElementReferenceException, NoSuchElementException])
numero_capture = 0


def capture(nom):
    global numero_capture
    numero_capture += 1
    fenetre.save_screenshot(str(arguments.captures / f"{numero_capture:02d}-{nom}.png"))


def bouton(texte):
    selecteur = f"//button[normalize-space()=\"{texte}\"]"
    return attente.until(lambda _: fenetre.find_element(By.XPATH, selecteur))


def cliquer(texte):
    # Svelte peut redessiner le bouton entre sa recherche et le clic : on recommence.
    def essayer(_):
        element = fenetre.find_element(By.XPATH, f"//button[normalize-space()=\"{texte}\"]")
        if not element.is_enabled():
            return False
        element.click()
        return True

    attente.until(essayer)


def aller_a_l_etape(etape):
    selecteur = f"nav button[data-etape='{etape}']"
    attente.until(lambda _: fenetre.find_element(By.CSS_SELECTOR, selecteur).click() or True)


def attendre_titre(titre):
    attente.until(lambda _: fenetre.find_element(By.TAG_NAME, "h1").text == titre)
    # Pas d'action en cours : la zone d'annonce est vide.
    attente.until(lambda _: not fenetre.find_elements(By.CSS_SELECTOR, "#action-en-cours span"))


def attendre_texte(texte):
    attente.until(lambda _: texte in fenetre.find_element(By.TAG_NAME, "main").text)


def remplir_dialogue(titre, chemin):
    # Dialogue GTK natif, retrouvé par son titre. Sans gestionnaire de fenêtres
    # (Xvfb), le focus doit lui être donné explicitement ; Ctrl+L ouvre la
    # saisie d'emplacement.
    recherche = subprocess.run(
        ["xdotool", "search", "--sync", "--onlyvisible", "--name", titre],
        check=True, capture_output=True, text=True, timeout=15,
    )
    dialogue = recherche.stdout.split()[0]
    # Le dialogue doit être entièrement affiché avant de recevoir le clavier.
    time.sleep(2)
    subprocess.run(["xdotool", "windowfocus", "--sync", dialogue], check=True)
    subprocess.run(["xdotool", "key", "--clearmodifiers", "ctrl+l"], check=True)
    time.sleep(0.3)
    subprocess.run(["xdotool", "type", "--delay", "10", str(chemin)], check=True)
    time.sleep(0.3)
    # Alt+O (« Ouvrir ») valide le chemin saisi, fichier comme dossier ; Entrée
    # serait d'abord consommée par l'autocomplétion.
    subprocess.run(["xdotool", "key", "alt+o"], check=True)


def verifier(condition, message):
    if not condition:
        raise AssertionError(message)
    print("OK", message)


try:
    # Accueil
    attendre_titre("Harmoniser les noms de swimlanes d'un export Signavio")
    verifier(not bouton("Inventorier les lanes").is_enabled(), "création impossible sans source ni dossier")
    capture("accueil")

    cliquer("Choisir le fichier SGX")
    remplir_dialogue("Choisir un export Signavio", source)
    attendre_texte("export.sgx")
    cliquer("Choisir le dossier des traitements")
    remplir_dialogue("Dossier qui accueillera le traitement", parent)
    attendre_texte(str(parent))
    capture("accueil-choix")
    cliquer("Inventorier les lanes")

    # Source
    attendre_titre("Source")
    verifier(fenetre.switch_to.active_element.tag_name == "h1", "le focus est placé sur le titre de l'étape")
    chiffres = [element.text for element in fenetre.find_elements(By.CSS_SELECTOR, "[data-chiffre]")]
    verifier(chiffres == ["2", "4", "2"], f"inventaire : 2 modèles, 4 occurrences, 2 noms ({chiffres})")
    verifier(dossier.joinpath("traitement.json").exists(), "dossier de traitement créé")
    capture("source")

    # Tentative interrompue avant toute production (revue D08) : l'écran
    # Résultat reste consultable pour en donner le dossier, sans succès affiché.
    etape_resultat = "nav button[data-etape='resultat']"
    verifier(not fenetre.find_element(By.CSS_SELECTOR, etape_resultat).is_enabled(), "résultat inaccessible sans tentative")
    interrompue = dossier / "sorties" / "tentative-001"
    interrompue.mkdir(parents=True)
    fenetre.execute_script("window.dispatchEvent(new Event('focus'))")
    attente.until(lambda _: fenetre.find_element(By.CSS_SELECTOR, etape_resultat).is_enabled())
    aller_a_l_etape("resultat")
    attendre_titre("Résultat")
    attendre_texte("Une production a été interrompue")
    verifier("SGX produit" not in fenetre.find_element(By.TAG_NAME, "main").text, "tentative interrompue non présentée comme un résultat")
    capture("resultat-interrompu")
    interrompue.rmdir()
    fenetre.execute_script("window.dispatchEvent(new Event('focus'))")
    attente.until(lambda _: not fenetre.find_element(By.CSS_SELECTOR, etape_resultat).is_enabled())
    aller_a_l_etape("source")
    attendre_titre("Source")

    # Correspondances : d'abord une cellule invalide, puis un classeur valide.
    cliquer("Préparer les correspondances")
    attendre_titre("Correspondances")
    edition_correspondances = dossier / "edition" / "correspondance_swimlanes.xlsx"
    ecrire_classeur(edition_correspondances, "Correspondance", [["Nom actuel", "Nouveau nom"], ["Serv. achats", 123]])
    cliquer("Lire les correspondances")
    attendre_texte("colonne B (Nouveau nom) : texte attendu (valeur lue : nombre 123)")
    verifier(fenetre.find_element(By.CSS_SELECTOR, "[role=alert]").is_displayed(), "erreur annoncée (role=alert)")
    verifier(not fenetre.find_elements(By.XPATH, "//button[normalize-space()=\"Préparer l'analyse\"]"), "rien n'est adopté après une erreur")
    capture("correspondances-erreur")

    ecrire_classeur(
        edition_correspondances,
        "Correspondance",
        [["Nom actuel", "Nouveau nom"], ["Serv. achats", "Service achats"], ["Achat", "Achats"]],
    )
    cliquer("Lire les correspondances")
    attendre_texte("2 nom(s) avec une demande de renommage ont été lus.")
    attendre_texte("« Achat »")
    verifier(bouton("Enregistrer une copie").is_enabled(), "correspondances : copie du classeur proposée")
    # Revue D04 : l'analyse n'est pas accessible avant d'être préparée.
    verifier(not fenetre.find_element(By.CSS_SELECTOR, "nav button[data-etape='analyse']").is_enabled(), "analyse inaccessible avant sa préparation")
    capture("correspondances-lues")

    # Analyse
    cliquer("Préparer l'analyse")
    attendre_titre("Analyse")
    lignes = fenetre.find_elements(By.CSS_SELECTOR, "tbody tr")
    verifier(len(lignes) == 2, "analyse : 2 propositions (un modèle homonyme par ligne)")
    # Revue du parcours : le classeur d'analyse (qui a lui aussi une colonne
    # Validation) n'est plus proposé à l'ouverture.
    verifier(not fenetre.find_elements(By.XPATH, "//button[normalize-space()=\"Ouvrir l'analyse\"]"), "analyse : pas d'ouverture du classeur d'analyse")
    capture("analyse")

    # Décisions contradictoires (OUI puis NON sur la même proposition) : bloquant.
    cliquer("Passer aux décisions")
    attendre_titre("Décisions")
    verifier(bouton("Enregistrer une copie").is_enabled(), "voie « faire valider ailleurs » proposée")
    edition_decisions = dossier / "edition" / "validation_modifications.xlsx"
    classeur = load_workbook(edition_decisions)
    feuille = classeur["Analyse"]
    ligne_a = next(ligne for ligne in range(2, feuille.max_row + 1) if feuille.cell(ligne, 6).value == MODELE_A)
    copie_ligne = [feuille.cell(ligne_a, colonne).value for colonne in range(1, feuille.max_column + 1)]
    feuille.cell(ligne_a, 5).value = "OUI"
    copie_ligne[4] = "NON"
    feuille.append(copie_ligne)
    classeur.save(edition_decisions)
    cliquer("Lire le classeur de décision")
    attendre_texte("ont reçu à la fois OUI et NON")
    verifier("décisions contraires pour « Serv. achats »" in fenetre.find_element(By.TAG_NAME, "main").text, "contradiction : lignes citées")
    verifier(not bouton("Générer le SGX modifié").is_enabled(), "contradiction : génération impossible")
    resultats = [cellule.text for cellule in fenetre.find_elements(By.CSS_SELECTOR, "tbody td:nth-child(4)")]
    verifier(resultats.count("CONTRADICTOIRE") == 2, f"contradiction : deux lignes marquées dans le tableau ({resultats})")
    capture("decisions-contradictoires")

    # Décisions : un OUI, une valeur non reconnue.
    classeur = load_workbook(edition_decisions)
    feuille = classeur["Analyse"]
    feuille.delete_rows(feuille.max_row)
    for ligne in range(2, feuille.max_row + 1):
        modele = feuille.cell(ligne, 6).value
        feuille.cell(ligne, 5).value = "OUI" if modele == MODELE_A else "à voir"
    classeur.save(edition_decisions)
    cliquer("Lire le classeur de décision")
    attendre_texte("le recomptage en mémoire est conforme")
    attendre_texte("depuis le classeur de décision du traitement")
    attendre_texte("La génération appliquera 1 décision(s) admise(s), soit 2 occurrence(s).")
    verifier("Validation « À VOIR » non reconnue" in fenetre.find_element(By.TAG_NAME, "main").text, "motif de la ligne en attente affiché")
    capture("decisions-controlees")

    # Import d'un retour au mauvais format : refusé, les décisions lues restent.
    cliquer("Importer un classeur de décision")
    remplir_dialogue("Choisir un classeur Excel", edition_correspondances)
    attendre_texte("Classeur non adopté. Les décisions lues le")
    verifier(bouton("Générer le SGX modifié").is_enabled(), "import refusé : décisions précédentes toujours utilisables")
    capture("decisions-import-refuse")

    # Revue D02 : relire des correspondances identiques ne perd pas le contrôle.
    aller_a_l_etape("correspondances")
    attendre_titre("Correspondances")
    cliquer("Relire les correspondances")
    attendre_texte("2 nom(s) avec une demande de renommage ont été lus.")
    aller_a_l_etape("decisions")
    attendre_titre("Décisions")
    attendre_texte("le recomptage en mémoire est conforme")
    verifier(bouton("Générer le SGX modifié").is_enabled(), "relecture identique : génération toujours possible")
    admises = fenetre.find_element(By.CSS_SELECTOR, "[data-bilan='admises']").text
    verifier(admises == "1", f"relecture identique : 1 décision admise conservée ({admises!r})")

    # Thème sombre : bascule, capture, puis retour au thème clair.
    theme = lambda: fenetre.execute_script("return document.documentElement.dataset.theme")
    theme_initial = theme()
    cliquer("Thème sombre")
    verifier(theme() == "dark" and bouton("Thème sombre").get_attribute("aria-pressed") == "true", "bascule vers le thème sombre")
    capture("decisions-theme-sombre")
    cliquer("Thème sombre")
    verifier(theme() == "light" and theme_initial == "light", "retour au thème clair")

    # Double clic sur « Générer » : une seule tentative.
    generer = bouton("Générer le SGX modifié")
    fenetre.execute_script("arguments[0].click(); arguments[0].click();", generer)
    attendre_titre("Résultat")
    attendre_texte("SGX produit et vérifié")
    tentatives = sorted(chemin.name for chemin in (dossier / "sorties").iterdir())
    verifier(tentatives == ["tentative-001"], f"une seule tentative malgré le double clic ({tentatives})")
    sgx = dossier / "sorties" / "tentative-001" / "export_modifie.sgx"
    verifier(lanes_du_sgx(sgx, MODELE_A) == ["Service achats", "Service achats", "Direction"], "modèle A renommé")
    verifier(lanes_du_sgx(sgx, MODELE_B) == ["Serv. achats"], "modèle B inchangé (décision en attente)")
    capture("resultat")

    # Classeur modifié dans Excel après lecture : génération suspendue.
    classeur = load_workbook(edition_decisions)
    classeur["Analyse"].cell(3, 5).value = "OUI"
    classeur.save(edition_decisions)
    fenetre.execute_script("window.dispatchEvent(new Event('focus'))")
    aller_a_l_etape("decisions")
    attendre_texte("Le classeur de décision a changé depuis sa dernière lecture")
    etape = fenetre.find_element(By.XPATH, "//nav//button[@aria-current='step']").text
    verifier("à relire" in etape, f"l'étape signale le classeur à relire ({etape!r})")
    verifier(not bouton("Générer le SGX modifié").is_enabled(), "génération suspendue tant que le classeur n'est pas relu")
    # Revue D03 : les compteurs sont présentés comme ceux du dernier contrôle.
    controle = fenetre.find_element(By.CSS_SELECTOR, "[data-controle='a-actualiser'] h2").text
    verifier(controle == "Dernier contrôle — à actualiser", f"compteurs signalés à actualiser ({controle!r})")
    admises = fenetre.find_element(By.CSS_SELECTOR, "[data-bilan='admises']").text
    verifier("à actualiser" in admises, f"bilan latéral signalé à actualiser ({admises!r})")
    capture("decisions-a-relire")

    # Correspondances modifiées dans Excel : l'analyse est « à actualiser ».
    ecrire_classeur(
        edition_correspondances,
        "Correspondance",
        [["Nom actuel", "Nouveau nom"], ["Serv. achats", "Service des achats"], ["Achat", "Achats"]],
    )
    fenetre.execute_script("window.dispatchEvent(new Event('focus'))")
    aller_a_l_etape("analyse")
    attendre_titre("Analyse")
    attendre_texte("Dernière analyse — à actualiser")
    verifier(fenetre.find_elements(By.CSS_SELECTOR, "[data-analyse='a-actualiser']"), "analyse : tableau marqué à actualiser")
    etape = fenetre.find_element(By.XPATH, "//nav//button[@aria-current='step']").text
    verifier("à actualiser" in etape, f"étape Analyse à actualiser ({etape!r})")
    propositions = fenetre.find_element(By.CSS_SELECTOR, "[data-bilan='propositions']").text
    verifier("à actualiser" in propositions, f"bilan : propositions à actualiser ({propositions!r})")
    verifier(not bouton("Passer aux décisions").is_enabled(), "analyse à actualiser : passage aux décisions suspendu")
    capture("analyse-a-actualiser")
    aller_a_l_etape("resultat")
    attendre_titre("Résultat")
    attendre_texte("ce résultat correspond aux dernières entrées lues")

    # Fermeture puis reprise du traitement.
    cliquer("Fermer le traitement")
    attendre_titre("Harmoniser les noms de swimlanes d'un export Signavio")
    cliquer("Ouvrir un dossier de traitement")
    remplir_dialogue("Ouvrir un dossier de traitement", dossier)
    attendre_titre("Résultat")
    verifier("SGX produit et vérifié" in fenetre.find_element(By.TAG_NAME, "main").text, "reprise : résultat retrouvé")
    capture("reprise")
    print("Parcours complet réussi.")
except (AssertionError, WebDriverException):
    capture("echec")
    raise
finally:
    fenetre.quit()
    pilote.terminate()
    shutil.rmtree(espace, ignore_errors=True)
